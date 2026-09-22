from http.server import ThreadingHTTPServer, SimpleHTTPRequestHandler
from pathlib import Path
from urllib.parse import urlparse
from openpyxl import load_workbook
from datetime import datetime
from threading import Lock
import base64
import html
import json
import os
import re
import shutil
import sys
import tempfile
import unicodedata

ROOT = Path(__file__).resolve().parent
TEMPLATE_DATA_FILE = ROOT / "data" / "LOYS_Aktien_Challenge.xlsx"
WEB_DIR = ROOT / "web"
LOCK = Lock()

HOST = os.environ.get("HOST", "0.0.0.0")
PORT = int(os.environ.get("PORT", "10000"))
ADMIN_USER = os.environ.get("ADMIN_USER", "")
ADMIN_PASSWORD = os.environ.get("ADMIN_PASSWORD", "")

requested_data_dir = Path(os.environ.get("DATA_DIR", str(ROOT / "data")))
try:
    requested_data_dir.mkdir(parents=True, exist_ok=True)
    DATA_DIR = requested_data_dir.resolve()
except Exception:
    DATA_DIR = (ROOT / "data").resolve()
    DATA_DIR.mkdir(parents=True, exist_ok=True)

DATA_FILE = DATA_DIR / "LOYS_Aktien_Challenge.xlsx"
if DATA_FILE != TEMPLATE_DATA_FILE and not DATA_FILE.exists():
    shutil.copy2(TEMPLATE_DATA_FILE, DATA_FILE)

MAG7_TICKERS = {"AAPL", "MSFT", "NVDA", "GOOGL", "GOOG", "AMZN", "META", "TSLA"}


def normalize_name(value):
    value = unicodedata.normalize("NFKD", str(value or ""))
    value = "".join(ch for ch in value if not unicodedata.combining(ch))
    value = value.lower().replace("&", " and ")
    value = re.sub(r"[^a-z0-9]+", " ", value)
    return re.sub(r"\s+", " ", value).strip()


def normalize_ticker(value):
    value = str(value or "").strip().upper()
    return re.sub(r"\s+", " ", value)


def ticker_root(value):
    value = normalize_ticker(value)
    return re.split(r"[\s\.:/\-]+", value, maxsplit=1)[0] if value else ""


def is_mag7_name(value):
    n = normalize_name(value)
    if not n:
        return False
    exact = {"apple", "microsoft", "nvidia", "alphabet", "google", "amazon", "meta", "facebook", "tesla"}
    if n in exact:
        return True
    prefixes = (
        "apple inc", "apple incorporated", "apple computer",
        "microsoft corp", "microsoft corporation",
        "nvidia corp", "nvidia corporation",
        "alphabet inc", "google llc",
        "amazon com", "amazon inc",
        "meta platforms", "facebook inc",
        "tesla inc", "tesla motors",
    )
    return any(n.startswith(p) for p in prefixes)


def is_mag7_pick(pick):
    return ticker_root(pick.get("ticker")) in MAG7_TICKERS or is_mag7_name(pick.get("name"))


def clean_pick(raw):
    raw = raw if isinstance(raw, dict) else {}
    return {
        "name": str(raw.get("name", "")).strip(),
        "ticker": normalize_ticker(raw.get("ticker", "")),
    }


def validate_pick(pick, label):
    name = pick["name"]
    ticker = pick["ticker"]
    if not name:
        return f"Bitte den Aktiennamen fuer {label} eingeben."
    if len(name) > 120:
        return f"Der Aktienname fuer {label} ist zu lang."
    if not ticker:
        return f"Bitte den Ticker fuer {label} eingeben."
    if len(ticker) > 24 or not re.match(r"^[A-Z0-9][A-Z0-9 .:/\-]{0,23}$", ticker):
        return f"Bitte einen gueltigen Ticker fuer {label} eingeben."
    if is_mag7_pick(pick):
        return "Aktien der Magnificent Seven sind in dieser Challenge ausgeschlossen."
    return None


def read_json(handler):
    length = int(handler.headers.get("Content-Length", 0))
    raw = handler.rfile.read(length) if length else b"{}"
    return json.loads(raw.decode("utf-8"))


def send_json(handler, payload, status=200):
    body = json.dumps(payload, ensure_ascii=False).encode("utf-8")
    handler.send_response(status)
    handler.send_header("Content-Type", "application/json; charset=utf-8")
    handler.send_header("Content-Length", str(len(body)))
    handler.send_header("Cache-Control", "no-store")
    handler.end_headers()
    handler.wfile.write(body)


def send_html(handler, markup, status=200):
    body = markup.encode("utf-8")
    handler.send_response(status)
    handler.send_header("Content-Type", "text/html; charset=utf-8")
    handler.send_header("Content-Length", str(len(body)))
    handler.send_header("Cache-Control", "no-store")
    handler.end_headers()
    handler.wfile.write(body)


def participation_rows(wb):
    ws = wb["Teilnahmen"]
    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r and r[0]:
            rows.append({"id": r[0], "timestamp": r[1], "name": r[2], "email": r[3]})
    return rows


def save_workbook_atomic(wb, destination):
    fd, temp_name = tempfile.mkstemp(prefix="loys_challenge_", suffix=".xlsx", dir=str(destination.parent))
    os.close(fd)
    try:
        wb.save(temp_name)
        os.replace(temp_name, destination)
    finally:
        if os.path.exists(temp_name):
            os.unlink(temp_name)


def admin_configured():
    return bool(ADMIN_USER and ADMIN_PASSWORD)


def admin_authorized(handler):
    if not admin_configured():
        return False
    auth = handler.headers.get("Authorization", "")
    expected = "Basic " + base64.b64encode(
        f"{ADMIN_USER}:{ADMIN_PASSWORD}".encode("utf-8")
    ).decode("ascii")
    return auth == expected


def request_admin_login(handler):
    if not admin_configured():
        return send_html(
            handler,
            "<h2>Admin-Zugang noch nicht eingerichtet</h2>"
            "<p>Bitte in Render die Environment Variables ADMIN_USER und ADMIN_PASSWORD setzen und neu deployen.</p>",
            503,
        )
    body = b"Admin-Zugang erforderlich."
    handler.send_response(401)
    handler.send_header("WWW-Authenticate", 'Basic realm="LOYS Aktien-Challenge Admin"')
    handler.send_header("Content-Type", "text/plain; charset=utf-8")
    handler.send_header("Content-Length", str(len(body)))
    handler.end_headers()
    handler.wfile.write(body)


def admin_page(handler):
    if not admin_authorized(handler):
        return request_admin_login(handler)

    with LOCK:
        wb = load_workbook(DATA_FILE, read_only=True, data_only=False)
        rows = participation_rows(wb)
        wb.close()

    count = len(rows)
    last = rows[-1] if rows else None
    last_text = "Noch keine Teilnahme gespeichert."
    if last:
        last_text = f"Letzte Teilnahme: {html.escape(str(last['timestamp']))} - {html.escape(str(last['name']))}"

    page = f'''<!doctype html>
<html lang="de"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>LOYS Challenge Admin</title>
<style>
body{{font-family:system-ui,-apple-system,Segoe UI,Arial,sans-serif;background:#f4f8fa;color:#13202a;margin:0;padding:24px}}
.box{{max-width:620px;margin:40px auto;background:#fff;border:1px solid #dce5ea;border-radius:18px;padding:26px;box-shadow:0 12px 34px rgba(11,31,45,.08)}}
h1{{margin:0 0 8px;font-size:28px}}p{{color:#657783;line-height:1.5}}.kpi{{font-size:42px;font-weight:800;color:#136A97;margin:18px 0 0}}.label{{font-size:12px;color:#657783;text-transform:uppercase;font-weight:700;letter-spacing:.06em}}a{{display:inline-block;margin-top:20px;background:#136A97;color:#fff;text-decoration:none;font-weight:750;padding:13px 16px;border-radius:11px}}.note{{font-size:12px;margin-top:22px;padding:12px;background:#fff6e8;border:1px solid #f0d6a8;border-radius:10px}}
</style></head><body><div class="box">
<h1>LOYS Aktien-Challenge</h1><p>Interner Admin-Bereich</p>
<div class="label">Gespeicherte Teilnahmen</div><div class="kpi">{count}</div>
<p>{last_text}</p>
<a href="/admin/download">Aktuelle Excel herunterladen</a>
<div class="note"><strong>Hinweis:</strong> Auf einem kostenlosen Render-Webservice ist der lokale Dateispeicher nicht dauerhaft. Die Excel deshalb regelmaessig herunterladen.</div>
</div></body></html>'''
    return send_html(handler, page)


def download_excel(handler):
    if not admin_authorized(handler):
        return request_admin_login(handler)
    if not DATA_FILE.exists():
        return send_json(handler, {"error": "Excel-Datei nicht gefunden."}, 404)

    with LOCK:
        body = DATA_FILE.read_bytes()

    filename = "LOYS_Aktien_Challenge_" + datetime.now().strftime("%Y-%m-%d_%H-%M") + ".xlsx"
    handler.send_response(200)
    handler.send_header("Content-Type", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    handler.send_header("Content-Disposition", f'attachment; filename="{filename}"')
    handler.send_header("Content-Length", str(len(body)))
    handler.send_header("Cache-Control", "no-store")
    handler.end_headers()
    handler.wfile.write(body)


class Handler(SimpleHTTPRequestHandler):
    def translate_path(self, path):
        parsed = urlparse(path)
        rel = parsed.path.lstrip("/") or "index.html"
        target = (WEB_DIR / rel).resolve()
        if not str(target).startswith(str(WEB_DIR.resolve())):
            return str(WEB_DIR / "index.html")
        return str(target)

    def log_message(self, fmt, *args):
        sys.stdout.write("%s - %s\n" % (self.address_string(), fmt % args))

    def do_GET(self):
        path = urlparse(self.path).path
        if path == "/health":
            return send_json(self, {"ok": True})
        if path == "/api/rules":
            return send_json(self, {
                "mag7_excluded": True,
                "mag7": ["Apple", "Microsoft", "NVIDIA", "Alphabet", "Amazon", "Meta Platforms", "Tesla"],
                "free_text_entry": True,
            })
        if path in ("/admin", "/admin/"):
            return admin_page(self)
        if path == "/admin/download":
            return download_excel(self)
        return super().do_GET()

    def do_POST(self):
        path = urlparse(self.path).path
        if path != "/api/submit":
            return send_json(self, {"error": "Nicht gefunden"}, 404)

        try:
            payload = read_json(self)
        except Exception:
            return send_json(self, {"error": "Ungueltige Anfrage"}, 400)

        name = str(payload.get("name", "")).strip()
        email = str(payload.get("email", "")).strip().lower()
        longs = [clean_pick(x) for x in payload.get("longs", [])]
        down = clean_pick(payload.get("down", {}))
        consent = bool(payload.get("consent"))

        if not name or len(name) > 80:
            return send_json(self, {"error": "Bitte Name oder Kuerzel angeben."}, 400)
        if not email or not re.match(r"^[^@\s]+@[^@\s]+\.[^@\s]+$", email):
            return send_json(self, {"error": "Bitte gueltige E-Mail-Adresse angeben."}, 400)
        if len(longs) != 3:
            return send_json(self, {"error": "Bitte genau drei Aktien fuer steigende Kurse eingeben."}, 400)

        for i, pick in enumerate(longs, start=1):
            error = validate_pick(pick, f"Favorit {i}")
            if error:
                return send_json(self, {"error": error}, 400)
        error = validate_pick(down, "die Aktie mit erwarteten fallenden Kursen")
        if error:
            return send_json(self, {"error": error}, 400)

        all_picks = longs + [down]
        ticker_keys = [normalize_ticker(p["ticker"]) for p in all_picks]
        name_keys = [normalize_name(p["name"]) for p in all_picks]
        if len(set(ticker_keys)) != 4:
            return send_json(self, {"error": "Bitte vier unterschiedliche Aktien eingeben. Ein Ticker wurde mehrfach verwendet."}, 400)
        if len(set(name_keys)) != 4:
            return send_json(self, {"error": "Bitte vier unterschiedliche Aktien eingeben. Ein Aktienname wurde mehrfach verwendet."}, 400)
        if not consent:
            return send_json(self, {"error": "Bitte Teilnahmebedingungen bestaetigen."}, 400)

        with LOCK:
            wb = load_workbook(DATA_FILE)
            existing = participation_rows(wb)
            if any(str(p.get("email", "")).lower() == email for p in existing):
                wb.close()
                return send_json(self, {"error": "Diese E-Mail-Adresse hat bereits teilgenommen."}, 409)

            ws = wb["Teilnahmen"]
            next_no = max(ws.max_row, 1)
            pid = f"LC-{datetime.now().strftime('%Y%m%d')}-{next_no:04d}"
            now = datetime.now().strftime("%d.%m.%Y %H:%M:%S")

            ws.append([
                pid, now, name, email,
                longs[0]["name"], longs[0]["ticker"],
                longs[1]["name"], longs[1]["ticker"],
                longs[2]["name"], longs[2]["ticker"],
                down["name"], down["ticker"],
                "Ja", "Bestaetigt",
            ])
            save_workbook_atomic(wb, DATA_FILE)
            wb.close()

        return send_json(self, {"ok": True, "id": pid})


if __name__ == "__main__":
    os.chdir(WEB_DIR)
    print(f"LOYS Aktien-Challenge laeuft auf http://{HOST}:{PORT}")
    print(f"Excel-Datei: {DATA_FILE}")
    print("Admin: /admin")
    ThreadingHTTPServer((HOST, PORT), Handler).serve_forever()
