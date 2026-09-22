from http.server import ThreadingHTTPServer, SimpleHTTPRequestHandler
from pathlib import Path
from urllib.parse import urlparse
from openpyxl import load_workbook
from datetime import datetime
from threading import Lock
import base64
import hmac
import json
import os
import re
import shutil
import sys
import unicodedata

ROOT = Path(__file__).resolve().parent
WEB_DIR = ROOT / "web"

TEMPLATE_DATA_FILE = ROOT / "data" / "LOYS_Aktien_Challenge.xlsx"
DATA_DIR = Path(os.environ.get("DATA_DIR", str(ROOT / "data"))).resolve()
DATA_FILE = DATA_DIR / "LOYS_Aktien_Challenge.xlsx"

HOST = os.environ.get("HOST", "0.0.0.0")
PORT = int(os.environ.get("PORT", "10000"))

ADMIN_USER = os.environ.get("ADMIN_USER", "")
ADMIN_PASSWORD = os.environ.get("ADMIN_PASSWORD", "")

LOCK = Lock()

DATA_DIR.mkdir(parents=True, exist_ok=True)
if DATA_FILE != TEMPLATE_DATA_FILE and not DATA_FILE.exists() and TEMPLATE_DATA_FILE.exists():
    shutil.copy2(TEMPLATE_DATA_FILE, DATA_FILE)

MAG7_TICKERS = {"AAPL", "MSFT", "NVDA", "GOOGL", "GOOG", "AMZN", "META", "TSLA"}


def normalize_name(value):
    value = unicodedata.normalize("NFKD", str(value or ""))
    value = "".join(ch for ch in value if not unicodedata.combining(ch))
    value = value.lower().replace("&", " and ")
    value = re.sub(r"[^a-z0-9]+", " ", value)
    return re.sub(r"\s+", " ", value).strip()


def normalize_ticker(value):
    value = str(value or "").strip().upper()
    value = re.sub(r"\s+", " ", value)
    return value


def ticker_root(value):
    value = normalize_ticker(value)
    return re.split(r"[\s\.:/\-]+", value, maxsplit=1)[0] if value else ""


def is_mag7_name(value):
    n = normalize_name(value)
    if not n:
        return False

    exact = {
        "apple", "microsoft", "nvidia", "alphabet", "google",
        "amazon", "meta", "facebook", "tesla"
    }
    if n in exact:
        return True

    prefixes = (
        "apple inc", "apple incorporated", "apple computer",
        "microsoft corp", "microsoft corporation",
        "nvidia corp", "nvidia corporation",
        "alphabet inc", "google llc",
        "amazon com", "amazon inc",
        "meta platforms", "facebook inc",
        "tesla inc", "tesla motors",
    )
    return any(n.startswith(p) for p in prefixes)


def is_mag7_pick(pick):
    return ticker_root(pick.get("ticker")) in MAG7_TICKERS or is_mag7_name(pick.get("name"))


def clean_pick(raw):
    raw = raw if isinstance(raw, dict) else {}
    return {
        "name": str(raw.get("name", "")).strip(),
        "ticker": normalize_ticker(raw.get("ticker", "")),
    }


def validate_pick(pick, label):
    name = pick["name"]
    ticker = pick["ticker"]

    if not name:
        return f"Bitte den Aktiennamen für {label} eingeben."

    if len(name) > 120:
        return f"Der Aktienname für {label} ist zu lang."

    if not ticker:
        return f"Bitte den Ticker für {label} eingeben."

    if len(ticker) > 24 or not re.match(r"^[A-Z0-9][A-Z0-9 .:/\-]{0,23}$", ticker):
        return f"Bitte einen gültigen Ticker für {label} eingeben."

    if is_mag7_pick(pick):
        return "Aktien der Magnificent Seven sind in dieser Challenge ausgeschlossen."

    return None


def read_json(handler):
    length = int(handler.headers.get("Content-Length", 0))
    raw = handler.rfile.read(length) if length else b"{}"
    return json.loads(raw.decode("utf-8"))


def send_json(handler, payload, status=200):
    body = json.dumps(payload, ensure_ascii=False).encode("utf-8")
    handler.send_response(status)
    handler.send_header("Content-Type", "application/json; charset=utf-8")
    handler.send_header("Content-Length", str(len(body)))
    handler.send_header("Cache-Control", "no-store")
    handler.end_headers()
    handler.wfile.write(body)


def participation_rows(wb):
    if "Teilnahmen" not in wb.sheetnames:
        return []

    ws = wb["Teilnahmen"]
    rows = []

    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[0]:
            continue
        rows.append({
            "id": row[0],
            "email": row[3] if len(row) > 3 else "",
        })

    return rows


def admin_credentials_configured():
    return bool(ADMIN_USER and ADMIN_PASSWORD)


def admin_authorized(handler):
    if not admin_credentials_configured():
        return False

    auth = handler.headers.get("Authorization", "")
    if not auth.startswith("Basic "):
        return False

    try:
        decoded = base64.b64decode(auth[6:]).decode("utf-8")
        supplied_user, supplied_password = decoded.split(":", 1)
    except Exception:
        return False

    return (
        hmac.compare_digest(supplied_user, ADMIN_USER)
        and hmac.compare_digest(supplied_password, ADMIN_PASSWORD)
    )


def request_admin_login(handler):
    if not admin_credentials_configured():
        body = (
            "Admin-Zugang ist noch nicht konfiguriert. "
            "Bitte in Render ADMIN_USER und ADMIN_PASSWORD als Environment Variables setzen."
        ).encode("utf-8")
        handler.send_response(503)
        handler.send_header("Content-Type", "text/plain; charset=utf-8")
        handler.send_header("Content-Length", str(len(body)))
        handler.send_header("Cache-Control", "no-store")
        handler.end_headers()
        handler.wfile.write(body)
        return

    body = b"Admin-Zugang erforderlich."
    handler.send_response(401)
    handler.send_header("WWW-Authenticate", 'Basic realm="LOYS Aktien-Challenge Admin"')
    handler.send_header("Content-Type", "text/plain; charset=utf-8")
    handler.send_header("Content-Length", str(len(body)))
    handler.send_header("Cache-Control", "no-store")
    handler.end_headers()
    handler.wfile.write(body)


def send_admin_page(handler):
    if not DATA_FILE.exists():
        count = 0
        file_status = "Excel-Datei wurde noch nicht gefunden."
    else:
        try:
            with LOCK:
                wb = load_workbook(DATA_FILE, read_only=True)
                count = len(participation_rows(wb))
                wb.close()
            file_status = "Excel-Datei ist verfügbar."
        except Exception:
            count = "?"
            file_status = "Excel-Datei konnte nicht gelesen werden."

    html = f"""<!DOCTYPE html>
<html lang="de">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>LOYS Aktien-Challenge – Admin</title>
<style>
body {{
    margin: 0;
    font-family: Arial, sans-serif;
    background: #f5f8fa;
    color: #13202a;
}}
.wrap {{
    max-width: 620px;
    margin: 0 auto;
    padding: 28px 18px 60px;
}}
.card {{
    background: #fff;
    border-radius: 18px;
    padding: 28px;
    box-shadow: 0 12px 35px rgba(11,31,45,.08);
}}
h1 {{
    margin-top: 0;
    color: #0b1f2d;
    font-size: 28px;
}}
.count {{
    font-size: 48px;
    font-weight: 800;
    color: #136A97;
    margin: 18px 0 4px;
}}
.label {{
    color: #6b7b87;
    margin-bottom: 24px;
}}
.btn {{
    display: block;
    text-align: center;
    background: #136A97;
    color: white;
    text-decoration: none;
    padding: 15px 18px;
    border-radius: 12px;
    font-weight: 700;
    margin-top: 18px;
}}
.note {{
    margin-top: 22px;
    font-size: 13px;
    line-height: 1.5;
    color: #6b7b87;
}}
</style>
</head>
<body>
<div class="wrap">
<div class="card">
<h1>Aktien-Challenge Admin</h1>
<div class="count">{count}</div>
<div class="label">gespeicherte Teilnahmen</div>
<p>{file_status}</p>
<a class="btn" href="/admin/download">Aktuelle Excel herunterladen</a>
<p class="note">Die Datei enthält die aktuell auf diesem Render-Service gespeicherten Einträge.</p>
</div>
</div>
</body>
</html>"""

    body = html.encode("utf-8")
    handler.send_response(200)
    handler.send_header("Content-Type", "text/html; charset=utf-8")
    handler.send_header("Content-Length", str(len(body)))
    handler.send_header("Cache-Control", "no-store")
    handler.end_headers()
    handler.wfile.write(body)


def download_excel(handler):
    if not DATA_FILE.exists():
        return send_json(handler, {"error": "Excel-Datei nicht gefunden."}, 404)

    with LOCK:
        body = DATA_FILE.read_bytes()

    stamp = datetime.now().strftime("%Y-%m-%d_%H-%M")
    filename = f"LOYS_Aktien_Challenge_{stamp}.xlsx"

    handler.send_response(200)
    handler.send_header(
        "Content-Type",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    handler.send_header(
        "Content-Disposition",
        f'attachment; filename="{filename}"',
    )
    handler.send_header("Content-Length", str(len(body)))
    handler.send_header("Cache-Control", "no-store")
    handler.end_headers()
    handler.wfile.write(body)


class Handler(SimpleHTTPRequestHandler):
    def translate_path(self, path):
        parsed = urlparse(path)
        rel = parsed.path.lstrip("/") or "index.html"
        target = (WEB_DIR / rel).resolve()

        if not str(target).startswith(str(WEB_DIR.resolve())):
            return str(WEB_DIR / "index.html")

        return str(target)

    def log_message(self, fmt, *args):
        sys.stdout.write("%s - %s\n" % (self.address_string(), fmt % args))

    def do_GET(self):
        path = urlparse(self.path).path.rstrip("/") or "/"

        if path == "/health":
            return send_json(self, {
                "ok": True,
                "data_file_exists": DATA_FILE.exists(),
                "admin_configured": admin_credentials_configured(),
            })

        if path == "/api/rules":
            return send_json(self, {
                "mag7_excluded": True,
                "mag7": [
                    "Apple", "Microsoft", "NVIDIA", "Alphabet",
                    "Amazon", "Meta Platforms", "Tesla"
                ],
                "free_text_entry": True,
            })

        if path == "/admin":
            if not admin_authorized(self):
                return request_admin_login(self)
            return send_admin_page(self)

        if path == "/admin/download":
            if not admin_authorized(self):
                return request_admin_login(self)
            return download_excel(self)

        return super().do_GET()

    def do_POST(self):
        path = urlparse(self.path).path.rstrip("/") or "/"

        if path != "/api/submit":
            return send_json(self, {"error": "Nicht gefunden."}, 404)

        try:
            payload = read_json(self)
        except Exception:
            return send_json(self, {"error": "Ungültige Anfrage."}, 400)

        name = str(payload.get("name", "")).strip()
        email = str(payload.get("email", "")).strip().lower()
        longs = [clean_pick(x) for x in payload.get("longs", [])]
        down = clean_pick(payload.get("down", {}))
        consent = bool(payload.get("consent"))

        if not name or len(name) > 80:
            return send_json(self, {"error": "Bitte Name oder Kürzel angeben."}, 400)

        if not email or not re.match(r"^[^@\s]+@[^@\s]+\.[^@\s]+$", email):
            return send_json(self, {"error": "Bitte gültige E-Mail-Adresse angeben."}, 400)

        if len(longs) != 3:
            return send_json(
                self,
                {"error": "Bitte genau drei Aktien für steigende Kurse eingeben."},
                400,
            )

        for i, pick in enumerate(longs, start=1):
            error = validate_pick(pick, f"Favorit {i}")
            if error:
                return send_json(self, {"error": error}, 400)

        error = validate_pick(
            down,
            "die Aktie mit erwarteten fallenden Kursen",
        )
        if error:
            return send_json(self, {"error": error}, 400)

        all_picks = longs + [down]
        ticker_keys = [normalize_ticker(p["ticker"]) for p in all_picks]
        name_keys = [normalize_name(p["name"]) for p in all_picks]

        if len(set(ticker_keys)) != 4:
            return send_json(
                self,
                {"error": "Bitte vier unterschiedliche Aktien eingeben. Ein Ticker wurde mehrfach verwendet."},
                400,
            )

        if len(set(name_keys)) != 4:
            return send_json(
                self,
                {"error": "Bitte vier unterschiedliche Aktien eingeben. Ein Aktienname wurde mehrfach verwendet."},
                400,
            )

        if not consent:
            return send_json(
                self,
                {"error": "Bitte Teilnahmebedingungen bestätigen."},
                400,
            )

        if not DATA_FILE.exists():
            return send_json(
                self,
                {"error": "Die Excel-Datei ist auf dem Server nicht vorhanden."},
                500,
            )

        with LOCK:
            wb = load_workbook(DATA_FILE)

            existing = participation_rows(wb)
            if any(str(p.get("email", "")).lower() == email for p in existing):
                wb.close()
                return send_json(
                    self,
                    {"error": "Diese E-Mail-Adresse hat bereits teilgenommen."},
                    409,
                )

            if "Teilnahmen" not in wb.sheetnames:
                wb.close()
                return send_json(
                    self,
                    {"error": "Das Tabellenblatt 'Teilnahmen' fehlt in der Excel-Datei."},
                    500,
                )

            ws = wb["Teilnahmen"]
            sequence_no = max(ws.max_row, 1)
            participant_id = (
                f"LC-{datetime.now().strftime('%Y%m%d')}-{sequence_no:04d}"
            )
            now = datetime.now().strftime("%d.%m.%Y %H:%M:%S")

            ws.append([
                participant_id,
                now,
                name,
                email,
                longs[0]["name"],
                longs[0]["ticker"],
                longs[1]["name"],
                longs[1]["ticker"],
                longs[2]["name"],
                longs[2]["ticker"],
                down["name"],
                down["ticker"],
                "Ja",
                "Bestätigt",
            ])

            wb.save(DATA_FILE)
            wb.close()

        return send_json(self, {
            "ok": True,
            "id": participant_id,
        })


if __name__ == "__main__":
    if not WEB_DIR.exists():
        raise FileNotFoundError(f"Web-Ordner nicht gefunden: {WEB_DIR}")

    print(f"LOYS Aktien-Challenge läuft auf http://{HOST}:{PORT}")
    print(f"Excel-Datei: {DATA_FILE}")
    print(f"Admin konfiguriert: {'Ja' if admin_credentials_configured() else 'Nein'}")

    os.chdir(WEB_DIR)
    ThreadingHTTPServer((HOST, PORT), Handler).serve_forever()
